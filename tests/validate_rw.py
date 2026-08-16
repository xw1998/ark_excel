# -*- coding: utf-8 -*-
from openpyxl import load_workbook

def check(name, fn):
    try:
        wb = load_workbook('tests/out/%s.xlsx' % name)
        errs = fn(wb)
        print(('PASS ' if not errs else 'FAIL ') + name, errs or '')
    except Exception as ex:
        print('FAIL %s load: %s' % (name, ex))

check('rw_basic', lambda wb: [] if (
    wb['Data']['B2'].value == 99.9 and wb['Data']['A4'].value == '新插入行'
    and str(wb['Data']['B3'].value) == '=B2*2'
    and 'A5:B5' in str(wb['Data'].merged_cells.ranges)
    and wb['Data'].freeze_panes == 'A2') else ['content mismatch'])

check('rw_zl', lambda wb: [] if wb['Data']['B2'].value == 1.5 else ['B2'])

def multi(wb):
    e = []
    if wb['Sheet1']['A1'].value != '表一A1': e.append('s1A1')
    if wb['Sheet1']['A3'].value != 'sheet1新增': e.append('s1A3')
    if wb['Sheet2']['A1'].value != '表二A1': e.append('s2A1')
    if wb['Sheet2']['A2'].value != 22: e.append('s2A2')
    if wb['Sheet2'].freeze_panes != 'B2': e.append('s2freeze')
    return e
check('rw_multi', multi)

def addsheet(wb):
    e = []
    if wb.sheetnames != ['Sheet1', 'Sheet2', 'NewSheet']: e.append('names:' + str(wb.sheetnames))
    if wb['NewSheet']['A1'].value != '新表': e.append('newA1')
    if wb['NewSheet']['A2'].value != 3.14: e.append('newA2')
    return e
check('rw_addsheet', addsheet)

def custstyle(wb):
    ws = wb['Data']
    e = []
    if ws['A7'].value != '自定义样式行': e.append('A7')
    if ws['A8'].value != '内置样式行': e.append('A8')
    if ws['A7'].fill.fgColor.rgb != 'FFFFD966': e.append('fill:' + str(ws['A7'].fill.fgColor.rgb))
    return e
check('rw_custstyle', custstyle)

def cfmt(wb):
    ws = wb['Data']
    e = []
    if ws['B2'].value != 88: e.append('B2')
    cf = list(ws.conditional_formatting)
    if not cf: e.append('no cond fmt')
    return e
check('rw_condfmt', cfmt)

def valid(wb):
    ws = wb['Data']
    e = []
    dv = ws.data_validations.dataValidation
    if not dv: e.append('no dv')
    elif 'A10' not in str(dv[0].sqref): e.append('range')
    elif dv[0].formula1 != '"合格,不合格,待检"': e.append('list:' + str(dv[0].formula1))
    return e
check('rw_valid', valid)

def opx(wb):
    e = []
    ws = wb['Third']
    if ws['A1'].value != 'openpyxl生成': e.append('A1:' + str(ws['A1'].value))
    if ws['B1'].value != 123.456: e.append('B1')
    if ws['A5'].value != '写入第三方模板': e.append('A5:' + str(ws['A5'].value))
    if 'A2:B2' not in str(ws.merged_cells.ranges): e.append('merge')
    if wb['Extra']['A1'].value != '第二个表': e.append('extra')
    return e
check('rw_openpyxl', opx)
