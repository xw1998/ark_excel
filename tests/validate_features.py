#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""ark_excel 全特性测试校验: 两层验证
  第一层: openpyxl 语义校验(值/样式/结构断言)
  第二层: Excel COM 真实打开(Excel 严格模式, 判损坏)
用法: python validate_features.py [out目录] [--no-com]
"""
import os
import sys
import openpyxl

OUT = sys.argv[1] if len(sys.argv) > 1 and not sys.argv[1].startswith('-') else 'out'
USE_COM = '--no-com' not in sys.argv

results = []          # (文件, 断言, 通过?)

def check(f, name, cond):
    results.append((f, name, bool(cond)))

def L(path):
    return openpyxl.load_workbook(os.path.join(OUT, path))

# ---------- basic ----------
wb = L('basic.xlsx'); ws = wb['S1']
check('basic', 'str',   ws['A1'].value == '文本')
check('basic', 'int',   ws['A2'].value == 42)
check('basic', 'float', ws['A3'].value == 3.14159)
check('basic', 'pct',   ws['A4'].value == -0.125)
check('basic', 'bool',  ws['A5'].value in (True, 1))
check('basic', 'error', ws['A6'].value == '#DIV/0!')
check('basic', 'blank styled', ws['A7'].value is None)

# ---------- escape ----------
ws = L('escape.xlsx')['S1']
check('escape', 'xml chars', ws['A1'].value == "a<b>&c\"d'e")
check('escape', 'utf8',      ws['A2'].value == '中文内容 UTF-8')
check('escape', 'tab kept',  ws['A3'].value == 'tab\there')

# ---------- merge ----------
ws = L('merge.xlsx')['S1']
m = sorted(str(r) for r in ws.merged_cells.ranges)
check('merge', 'horizontal band', 'A1:F1' in m)
check('merge', 'two-level header', m.count('A2:C2') + m.count('D2:F2') == 2)
check('merge', 'row span', 'C4:C6' in m)
check('merge', 'block 3x3', 'D3:F5' in m)
check('merge', 'count', len(m) == 6)

# ---------- merge_stagger ----------
ws = L('merge_stagger.xlsx')['S1']
m = sorted(str(r) for r in ws.merged_cells.ranges)
check('merge_stagger', 'pair left',  'A2:B2' in m)
check('merge_stagger', 'pair right', 'C2:D2' in m)
check('merge_stagger', 'offset mid', 'B3:C3' in m)
check('merge_stagger', 'resume A',   'A4:B4' in m)
check('merge_stagger', 'resume B',   'C4:D4' in m)
check('merge_stagger', 'no overlap', len(m) == 5)
check('merge_stagger', 'value in topleft', ws['A2'].value == '左组')
check('merge_stagger', 'data below', ws['A5'].value == 1)

# ---------- merge_rowspan(双行合并+错开) ----------
ws = L('merge_rowspan.xlsx')['S1']
m = sorted(str(r) for r in ws.merged_cells.ranges)
check('merge_rowspan', 'vertical 2-row', 'A2:A3' in m)
check('merge_rowspan', 'block 2x2',      'C2:D3' in m)
check('merge_rowspan', 'offset 2-row',    'B4:B5' in m)
check('merge_rowspan', 'offset 2x2',      'C4:D5' in m)
check('merge_rowspan', 'no overlap',      len(m) == 4)
check('merge_rowspan', 'value in topleft', ws['A2'].value == '竖向双行')
check('merge_rowspan', 'offset value',     ws['B4'].value == '错开双行')
check('merge_rowspan', 'data below',       ws['A6'].value == 42)

# ---------- altrow ----------
ws = L('altrow.xlsx')['S1']
check('altrow', 'odd alt fill', ws['A3'].fill.fgColor.rgb in ('00F2F2F2', 'FFF2F2F2'))
check('altrow', 'even plain',   ws['A4'].fill.patternType is None)
check('altrow', 'values', ws['B7'].value == 7 * 1.5)

# ---------- date ----------
ws = L('date.xlsx')['S1']
check('date', 'is_date', ws['A1'].is_date)
check('date', '2024-03-15', str(ws['A1'].value)[:10] == '2024-03-15')
from datetime import date as _date
# 1900-02-28 的 Excel 序列号应为 59(1900"闰日"bug 由天数>60 时+1 补偿)
check('date', '1900 leap bug handled',
      ws['A2'].value.toordinal() - _date(1899, 12, 31).toordinal() == 59)
check('date', '1900-03-01 serial 61',
      ws['A3'].value.toordinal() - _date(1899, 12, 31).toordinal() == 60)
check('date', 'datetime has time', str(ws['A5'].value) >= '2026-08-16 14:30:05'
      if ws['A5'].is_date else False)

# ---------- formula ----------
ws = L('formula.xlsx')['S1']
check('formula', 'add',    ws['A3'].value == '=A1+A2')
check('formula', 'sum',    ws['A4'].value == '=SUM(A1:A2)*2')
check('formula', 'if str', ws['A5'].value == '=IF(A1>5,"big","small")')

# ---------- freeze ----------
ws = L('freeze.xlsx')['S1']
check('freeze', 'pane B2', str(ws.freeze_panes) == 'B2')
check('freeze', 'autofilter', str(ws.auto_filter.ref) == 'A1:B20')

# ---------- colrow ----------
ws = L('colrow.xlsx')['S1']
check('colrow', 'width A', abs(ws.column_dimensions['A'].width - 30.5) < 0.2)
check('colrow', 'width B', abs(ws.column_dimensions['B'].width - 8.25) < 0.2)
check('colrow', 'height r1', abs(ws.row_dimensions[1].height - 40) < 1.0)
check('colrow', 'height r2', abs(ws.row_dimensions[2].height - 6.5) < 0.4)

# ---------- hidden ----------
ws = L('hidden.xlsx')['S1']
check('hidden', 'row3', ws.row_dimensions[3].hidden)
check('hidden', 'col B', ws.column_dimensions['B'].hidden)
check('hidden', 'row1 not hidden', not ws.row_dimensions[1].hidden)

# ---------- styles ----------
wb = L('styles.xlsx'); ws = wb['S1']
c = ws['A1']
check('styles', 'bold 16 red', c.font.b and c.font.sz == 16 and
      c.font.color.rgb in ('00FF0000', 'FFFF0000'))
check('styles', 'thick border', c.border.left.style == 'thick')
check('styles', 'custom fmt', ws['A3'].number_format == '0.000"kg"')
check('styles', 'builtin pct', ws['A2'].number_format == '0%')
check('styles', 'money fmt', ws['A7'].number_format == '#,##0.00')

# ---------- condfmt ----------
ws = L('condfmt.xlsx')['S1']
cfs = list(ws.conditional_formatting)
check('condfmt', 'two ranges rules', len(cfs) >= 1 and len(cfs[0].rules) == 2)
ops = {r.operator for r in cfs[0].rules}
check('condfmt', 'operators', ops == {'greaterThan', 'lessThan'})
check('condfmt', 'formula', cfs[0].rules[0].formula in (['50'], ['20']))

# ---------- valid ----------
ws = L('valid.xlsx')['S1']
dvs = ws.data_validations.dataValidation
check('valid', 'two rules', len(dvs) == 2)
lst = [v for v in dvs if v.type == 'list']
rng = [v for v in dvs if v.type == 'whole']
check('valid', 'list formula', lst and '告警' in lst[0].formula1)
check('valid', 'whole between', rng and rng[0].operator == 'between'
      and rng[0].formula1 == '0' and rng[0].formula2 == '100')
check('valid', 'prompt', lst and lst[0].prompt == '请选择状态')

# ---------- hyperlink ----------
wb = L('hyperlink.xlsx'); ws = wb['S1']
check('hyperlink', 'url', ws['A1'].hyperlink.target ==
      'https://gitee.com/xw19981010/svcrtos_new.git')
check('hyperlink', 'url with params', ws['A2'].hyperlink.target ==
      'https://e.com/?a=1&b=2')
check('hyperlink', 'internal', ws['A3'].hyperlink.location == '目标页!A1')
check('hyperlink', 'target sheet exists', '目标页' in wb.sheetnames)

# ---------- names ----------
wb = L('names.xlsx')
check('names', 'defined', set(['DataZone', 'Answer']) <= set(wb.defined_names))
check('names', 'refers', wb.defined_names['DataZone'].attr_text == "S1!$A$1:$A$2")

# ---------- multisheet ----------
wb = L('multisheet.xlsx')
check('multisheet', '5 sheets', wb.sheetnames == ['S1', '表2', '表3', '表4', '表5'])
check('multisheet', 'tab color', wb['表2'].sheet_properties.tabColor is not None)
check('multisheet', 'per-sheet cells', wb['表5']['A2'].value == 4 * 1.25)

# ---------- print ----------
wb = L('print.xlsx'); ws = wb['S1']
check('print', 'landscape A3', ws.page_setup.orientation == 'landscape'
      and int(ws.page_setup.paperSize) == 8)
check('print', 'print area', 'A1:A25' in str(ws.print_area).replace('$', ''))
hf = ws.oddHeader.left.text if ws.oddHeader else None
check('print', 'header', hf == '左')
check('print', 'row break', any(b.id == 15 for b in ws.row_breaks.brk))

# ---------- view ----------
ws = L('view.xlsx')['S1']
check('view', 'rtl', ws.sheet_view.rightToLeft)
check('view', 'gridlines off', ws.sheet_view.showGridLines in (False, None, '0'))
check('view', 'tab color', ws.sheet_properties.tabColor is not None)

# ---------- sst ----------
ws = L('sst.xlsx')['S1']
check('sst', 'unique strings', ws['A60'].value == '字符串条目编号060号')
check('sst', 'dedup', ws['B20'].value == '字符串条目编号001号')

# ---------- overwrite ----------
ws = L('overwrite.xlsx')['S1']
check('overwrite', 'later write wins', ws['C10'].value == '覆盖')
check('overwrite', 'odd col kept', ws['A5'].value == 105)
check('overwrite', 'sorted rows', ws.max_row == 20)

# ---------- 汇总 ----------
fails = [(f, n) for f, n, c in results if not c]
total = len(results)
print('openpyxl: %d/%d passed' % (total - len(fails), total))
for f, n in fails:
    print('  FAIL', f, '-', n)

# ---------- 第二层: Excel COM ----------
com_fail = []
if USE_COM:
    try:
        import pythoncom
        import win32com.client
        pythoncom.CoInitialize()
        xl = win32com.client.DispatchEx('Excel.Application')
        xl.Visible = False
        xl.DisplayAlerts = False
        for fn in sorted(os.listdir(OUT)):
            if not fn.endswith('.xlsx'):
                continue
            p = os.path.abspath(os.path.join(OUT, fn))
            try:
                w = xl.Workbooks.Open(p, 0, True)
                w.Close(False)
                print('COM PASS', fn)
            except Exception:
                com_fail.append(fn)
                print('COM FAIL', fn)
        xl.Quit()
    except ImportError:
        print('pywin32 不可用, 跳过 COM 验证')

ok = not fails and not com_fail
print('====', 'ALL OK' if ok else 'SOME FAILED',
      '| openpyxl %d/%d | COM fail %d' % (total - len(fails), total, len(com_fail)))
sys.exit(0 if ok else 1)
