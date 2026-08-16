#!/usr/bin/env python3
"""ark_excel 自动化验证: 用 openpyxl 解开生成的 xlsx 并核对各项特性"""
import sys
import openpyxl

path = sys.argv[1] if len(sys.argv) > 1 else 'demo_report.xlsx'
wb = openpyxl.load_workbook(path)
ws = wb['设备巡检']
ok = []

def check(name, cond):
    ok.append((name, bool(cond)))
    print(('PASS' if cond else 'FAIL'), name)

check('sheets', wb.sheetnames == ['设备巡检', 'Sheet2'])
check('merged', sorted(str(m) for m in ws.merged_cells.ranges) ==
      ['A11:B11', 'A1:F1', 'A4:B4', 'C4:F4'])
check('title', ws['A1'].value == '车间设备巡检日报表')
check('date', ws['B2'].is_date and ws['B2'].value is not None)
check('formula', ws['C11'].value == '=AVERAGE(C6:C10)')
check('freeze', str(ws.freeze_panes) == 'A6')
check('autofilter', str(ws.auto_filter.ref) == 'A5:F10')
check('colwidth', abs(ws.column_dimensions['A'].width - 16) < 0.1)
check('rowheight', abs(ws.row_dimensions[1].height - 24) < 0.6)
check('num cell', ws['C6'].value == 25.5)
check('int cell', ws['D6'].value == 100)
check('error cell', ws['F7'].value == '#N/A')
check('hyperlink internal', ws['A12'].hyperlink is not None and
      ws['A12'].hyperlink.location == 'Sheet2!A1')
check('hyperlink url', ws['C12'].hyperlink is not None and
      ws['C12'].hyperlink.target == 'https://example.com/spec.pdf')
dv = ws.data_validations.dataValidation
check('validation list', dv and dv[0].type == 'list' and
      '正常' in dv[0].formula1 and str(dv[0].sqref) == 'E6:E20')
cf = list(ws.conditional_formatting)
rule = cf[0].rules[0] if cf else None
check('condfmt', rule is not None and rule.operator == 'greaterThan' and
      rule.formula == ['130'])
check('page landscape', ws.page_setup.orientation == 'landscape')
check('print area', str(ws.print_area).replace('$','').replace("'",'') == '设备巡检!A1:F12')
check('defined name', 'TrendData' in wb.defined_names)
check('tab color', wb['Sheet2'].sheet_properties.tabColor is not None)
check('custom style', ws['E7'].fill.fgColor.rgb in ('00FFF2CC','FFFFF2CC'))
check('alt row', ws['A7'].fill.fgColor.rgb in ('00F2F2F2', 'FFF2F2F2'))
check('header style bg', ws['A5'].fill.fgColor.rgb in ('004472C4', 'FF4472C4'))
check('header font white', ws['A5'].font.color.rgb in ('00FFFFFF', 'FFFFFFFF'))

fails = [n for n, c in ok if not c]
print('---- %d/%d passed' % (len(ok) - len(fails), len(ok)))
sys.exit(1 if fails else 0)
